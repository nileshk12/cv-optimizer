from fastapi import APIRouter
from pydantic import BaseModel
import spacy
from spacy.matcher import PhraseMatcher
from docx import Document
import os
import uuid
from fastapi.responses import FileResponse
import time
from skillNer.skill_extractor_class import SkillExtractor
from skillNer.general_params import SKILL_DB
from sklearn.feature_extraction.text import TfidfVectorizer
import numpy as np
from typing import List

router = APIRouter()

# Load spaCy model and initialize SkillExtractor
nlp = spacy.load("en_core_web_sm")
matcher = PhraseMatcher(nlp.vocab)
skill_extractor = SkillExtractor(nlp, SKILL_DB, matcher)

# Directory to store optimized CVs
OUTPUT_DIR = "static/optimized_cvs"
os.makedirs(OUTPUT_DIR, exist_ok=True)

class JobDescription(BaseModel):
    title: str
    description: str
    resume_keywords: List[str]
    resume_content: str

def extract_keywords(text: str, use_tfidf: bool = False) -> List[str]:
    # Annotate text with skillNer to extract skills
    annotations = skill_extractor.annotate(text)
    
    # Extract skills from annotations
    skills = []
    for match in annotations["results"]["full_matches"]:
        skills.append(match["doc_node_value"].lower())
    for match in annotations["results"]["ngram_scored"]:
        skills.append(match["doc_node_value"].lower())
    
    skills = sorted(list(set(skills)))
    
    if use_tfidf and skills:
        # Use TF-IDF to rank skills by importance in the JD
        documents = [text] + skills
        vectorizer = TfidfVectorizer()
        tfidf_matrix = vectorizer.fit_transform(documents)
        feature_names = vectorizer.get_feature_names_out()
        
        tfidf_scores = tfidf_matrix[0].toarray().flatten()
        skill_scores = {}
        for skill in skills:
            if skill in feature_names:
                idx = list(feature_names).index(skill)
                skill_scores[skill] = tfidf_scores[idx]
        
        skills = sorted(skill_scores.keys(), key=lambda x: skill_scores[x], reverse=True)
    
    return skills

def cleanup_old_files():
    current_time = time.time()
    for filename in os.listdir(OUTPUT_DIR):
        filepath = os.path.join(OUTPUT_DIR, filename)
        if os.path.isfile(filepath):
            file_age = current_time - os.path.getmtime(filepath)
            if file_age > 3600:  # 1 hour
                os.remove(filepath)

def create_optimized_cv(resume_content: str, missing_skills: List[str]) -> str:
    cleanup_old_files()
    
    doc = Document()
    lines = resume_content.split("\n")
    skills_section_found = False
    experience_section_found = False
    
    for line in lines:
        if "experience" in line.lower():
            experience_section_found = True
            doc.add_paragraph(line)
            continue
        if experience_section_found and line.strip().startswith("-"):
            doc.add_paragraph(line)
        else:
            doc.add_paragraph(line)
        if "skills" in line.lower():
            skills_section_found = True
    
    if not skills_section_found:
        doc.add_heading("Skills", level=2)
    if missing_skills:
        skills_text = ", ".join(missing_skills)
        doc.add_paragraph(f"Added Skills (based on JD): {skills_text}")
    
    filename = f"optimized_cv_{uuid.uuid4()}.docx"
    filepath = os.path.join(OUTPUT_DIR, filename)
    doc.save(filepath)
    
    return filename

@router.post("/analyze")
async def analyze_job(job: JobDescription):
    # Extract JD keywords with TF-IDF ranking
    jd_keywords = extract_keywords(job.description, use_tfidf=True)
    missing_skills = [keyword for keyword in jd_keywords if keyword not in job.resume_keywords]
    
    suggestions = []
    if missing_skills:
        suggestions.append(f"Added the following skills to your resume: {', '.join(missing_skills)}.")
    else:
        suggestions.append("Your resume already covers most of the job description keywords!")
    
    optimized_cv_filename = create_optimized_cv(job.resume_content, missing_skills)
    download_url = f"http://54.212.243.217:8000/static/optimized_cvs/{optimized_cv_filename}"
    
    return {
        "title": job.title,
        "jd_keywords": jd_keywords,
        "missing_skills": missing_skills,
        "suggestions": suggestions,
        "optimized_cv_url": download_url
    }

@router.get("/download/{filename}")
async def download_optimized_cv(filename: str):
    filepath = os.path.join(OUTPUT_DIR, filename)
    if os.path.exists(filepath):
        return FileResponse(filepath, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", filename=filename)
    return {"error": "File not found"}
